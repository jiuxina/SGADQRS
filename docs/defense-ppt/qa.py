# -*- coding: utf-8 -*-
"""Programmatic QA for generated PPTX: bounds, text overflow estimate, overlap."""
import sys, math
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
SLIDE_W = 13.3333
SLIDE_H = 7.5

def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh

def text_of(sh):
    if not sh.has_text_frame:
        return ""
    return "".join(p.text for p in sh.text_frame.paragraphs)

def main(path):
    prs = Presentation(path)
    issues = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in iter_shapes(slide.shapes):
            try:
                left, top, w, h = sh.left, sh.top, sh.width, sh.height
            except Exception:
                continue
            if left is None or w is None:
                continue
            x0, y0 = left / EMU_IN, top / EMU_IN
            x1, y1 = (left + w) / EMU_IN, (top + h) / EMU_IN
            # 1) out of slide bounds
            if x0 < -0.02 or y0 < -0.02 or x1 > SLIDE_W + 0.02 or y1 > SLIDE_H + 0.02:
                issues.append((si, "OUT-OF-SLIDE", f"({x0:.2f},{y0:.2f},{x1:.2f},{y1:.2f})", sh.name or text_of(sh)[:20]))
            # 2) text overflow estimate (CJK full-width ~ 1em per char)
            txt = text_of(sh)
            if not txt or not sh.has_text_frame:
                continue
            max_pt = 0
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        max_pt = max(max_pt, r.font.size.pt)
            if max_pt == 0:
                continue
            font_pt = max_pt
            # available chars per line (approximate: CJK glyph width = 1.0 * font size; latin ~0.55)
            usable_w = (w / EMU_IN) * 72 - 4  # points, minus margins
            cjk = sum(1 for c in txt if ord(c) > 0x2E80)
            latin = len(txt) - cjk
            line_w_pt = cjk * font_pt + latin * font_pt * 0.55
            lines_needed = max(1, math.ceil(line_w_pt / max(usable_w, 10)))
            line_h_pt = font_pt * 1.42
            need_h = lines_needed * line_h_pt
            box_h_pt = (h / EMU_IN) * 72
            if need_h > box_h_pt * 1.12:
                issues.append((si, "TEXT-OVERFLOW", f"need~{need_h:.0f}pt in {box_h_pt:.0f}pt @{font_pt}pt", txt[:26]))
    # 3) pairwise text-box overlap (same slide)
    boxes = {}
    for si, slide in enumerate(prs.slides, 1):
        arr = []
        for sh in iter_shapes(slide.shapes):
            if not sh.has_text_frame or not text_of(sh):
                continue
            try:
                arr.append((sh.left, sh.top, sh.width, sh.height, text_of(sh)[:18]))
            except Exception:
                pass
        boxes[si] = arr
    for si, arr in boxes.items():
        for i in range(len(arr)):
            for j in range(i + 1, len(arr)):
                a, b = arr[i], arr[j]
                ax0, ay0, ax1, ay1 = a[0]/EMU_IN, a[1]/EMU_IN, (a[0]+a[2])/EMU_IN, (a[1]+a[3])/EMU_IN
                bx0, by0, bx1, by1 = b[0]/EMU_IN, b[1]/EMU_IN, (b[0]+b[2])/EMU_IN, (b[1]+b[3])/EMU_IN
                ox = min(ax1, bx1) - max(ax0, bx0)
                oy = min(ay1, by1) - max(ay0, by0)
                if ox > 0.05 and oy > 0.05:
                    issues.append((si, "OVERLAP", f"({a[4]!r}) x ({b[4]!r})", f"overlap {ox:.2f}x{oy:.2f}in"))
    if not issues:
        print("QA PASS: no issues found")
    else:
        print(f"QA REPORT: {len(issues)} issue(s)")
        for it in issues:
            print(f"  slide {it[0]:>2} | {it[1]:<14} | {it[3]}")
        sys.exit(1)

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "赛友TeamUp-答辩PPT.pptx")
